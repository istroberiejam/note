import base64,io,json,sys,threading,unittest,urllib.request
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from server import Handler,ThreadingHTTPServer,extract
from docx import Document
from pptx import Presentation
class Workflow(unittest.TestCase):
 @classmethod
 def setUpClass(cls):
  cls.server=ThreadingHTTPServer(('127.0.0.1',0),Handler);cls.url='http://127.0.0.1:'+str(cls.server.server_port);threading.Thread(target=cls.server.serve_forever,daemon=True).start()
 @classmethod
 def tearDownClass(cls):cls.server.shutdown();cls.server.server_close()
 def post(self,path,p):
  req=urllib.request.Request(self.url+'/api/'+path,data=json.dumps(p).encode(),headers={'Content-Type':'application/json'})
  with urllib.request.urlopen(req) as r:return json.load(r)
 def test_docx_tables_preserved(self):
  d=Document();d.add_paragraph('Concept');d.add_table(rows=1,cols=1).cell(0,0).text='Exact table text';b=io.BytesIO();d.save(b)
  result=self.post('extract',{'name':'sample.docx','data':base64.b64encode(b.getvalue()).decode()})
  self.assertEqual(result['pages'][0]['text'],'Concept\nExact table text')
 def test_slides(self):
  p=Presentation();s=p.slides.add_slide(p.slide_layouts[1]);s.shapes.title.text='Accuracy';s.placeholders[1].text='Closeness to true value';b=io.BytesIO();p.save(b)
  pages=extract('sample.pptx',b.getvalue());self.assertEqual(pages[0]['location'],'Slide 1');self.assertIn('Closeness to true value',pages[0]['text'])
 def test_export_sections(self):
  p={'title':'Chemistry','concepts':[{'title':'Accuracy','text':'Exact source','source':'Slide 1','side':'Teacher discussion'}],'references':[{'name':'Book chapter','text':'Reference text'}],'quiz':'1. Answer','exam':'1. A, eliminate B'}
  r=self.post('export',p);d=Document(io.BytesIO(base64.b64decode(r['data'])));text='\n'.join(p.text for p in d.paragraphs)
  for expected in ['Exact source','Teacher discussion','Book references','Reference text','Quiz review','Exam review']:self.assertIn(expected,text)
 def test_static_ui(self):
  with urllib.request.urlopen(self.url) as r:self.assertIn(b'Class notebook',r.read())
if __name__=='__main__':unittest.main()
