import io,json,unittest
from docx import Document
from docx.oxml.ns import qn
import exports as e
PROJECT={'title':'Clinical Chemistry 🧪','metadata':{'courseCode':'CC1','professor':'Dr. Reyes','module':'1: Introduction','semester':'First semester','student':'Student · Section A'},'concepts':[{'title':'Accuracy','text':'Exact source wording.\nSecond original line.','source':'slides.pptx · Slide 2','side':'Teacher highlights the true value.'}],'references':[{'name':'Reference chapter','location':'Page 22','related':'Accuracy','text':'Book reference wording.'}],'quiz':'1. Identify accuracy.\nAnswer: Closeness to the true value.','exam':'1. Which is correct?\nA. Accuracy\nB. Precision\nAnswer: A.\nElimination: B concerns repeatability.'}
class TemplateExportTests(unittest.TestCase):
 def test_word_page_columns_type_and_side_box(self):
  d=Document(io.BytesIO(e.word(PROJECT)));text='\n'.join([p.text for p in d.paragraphs]+[c.text for t in d.tables for row in t.rows for c in row.cells])
  for value in ['TOPIC OVERVIEW','I. Accuracy','Exact source wording.\nSecond original line.','SIDE NOTE/S','Teacher highlights','CC1','Dr. Reyes','Page 22','Related concept: Accuracy','Quiz review','Exam review']:self.assertIn(value,text)
  self.assertAlmostEqual(d.sections[0].page_width.mm,210,places=1)
  self.assertAlmostEqual(d.sections[0].top_margin.mm,12.7,places=1)
  self.assertEqual([s._sectPr.find(qn('w:cols')).get(qn('w:num')) for s in d.sections],['1','2']*4)
  title=d.paragraphs[0].runs[0];self.assertEqual(title.font.name,'Poppins Black');self.assertEqual(title.font.size.pt,34)
  banner=next(p for p in d.paragraphs if p.text=='TOPIC OVERVIEW');self.assertEqual(banner._p.find('.//'+qn('w:shd')).get(qn('w:fill')),e.BLUE)
  side=next(p for p in d.paragraphs if p.text.startswith('Teacher highlights'));self.assertIsNotNone(side._p.find('.//'+qn('w:pBdr')))
  import zipfile
  with zipfile.ZipFile(io.BytesIO(e.word(PROJECT))) as z:self.assertIn('word/fonts/Poppins-bold.ttf',z.namelist());self.assertNotIn('word/fonts/HelveticaNeue-regular.ttf',z.namelist())
  self.assertNotIn('Input definition here',text);self.assertNotIn('CreativiJI',text)
 def test_google_utf16_tab_ranges_and_column_sections(self):
  plan=e.google_plan(PROJECT,'notes','notes-tab');text=plan[0]['insertText']['text'];self.assertIn('Exact source wording.\nSecond original line.',text)
  break_index=plan[1]['insertSectionBreak']['location']['index']
  self.assertEqual(break_index,1+e.units(text[:text.index('TOPIC OVERVIEW')]))
  section=[r['updateSectionStyle'] for r in plan if 'updateSectionStyle' in r]
  self.assertEqual(len(section[0]['sectionStyle']['columnProperties']),1);self.assertEqual(len(section[1]['sectionStyle']['columnProperties']),2)
  banners=[r['updateParagraphStyle'] for r in plan if 'updateParagraphStyle' in r and 'shading' in r['updateParagraphStyle']['paragraphStyle']]
  first=banners[0]['range'];self.assertEqual(first['startIndex'],break_index+2)
  self.assertEqual(first['endIndex']-first['startIndex'],e.units('TOPIC OVERVIEW\n'))
  for req in plan:
   for body in req.values():
    if 'range' in body:self.assertEqual(body['range']['tabId'],'notes-tab');self.assertLessEqual(body['range']['endIndex'],1+e.units(text)+2)
  self.assertTrue(any('borderLeft' in r.get('updateParagraphStyle',{}).get('paragraphStyle',{}) for r in plan))
 def test_google_reuses_first_tab_and_formats_all_four(self):
  calls=[]
  def call(url,p,token):
   calls.append((url,p));self.assertEqual(token,'test-token')
   if not url.endswith(':batchUpdate'):return {'documentId':'new-doc','tabs':[{'tabProperties':{'tabId':'initial'}}]}
   if 'addDocumentTab' in p['requests'][-1]:return {'replies':[{},*({'addDocumentTab':{'tabProperties':{'tabId':id}}} for id in ['ref','quiz','exam'])]}
   return {}
  result=e.google_export(PROJECT,'test-token',call,lambda *a:None);self.assertIn('new-doc',result['url'])
  self.assertEqual(calls[1][1]['requests'][0]['updateDocumentTabProperties']['tabProperties']['title'],'Class notes')
  insertions=[r['insertText'] for _,body in calls[2:] for r in body['requests'] if 'insertText' in r]
  self.assertEqual([r['endOfSegmentLocation']['tabId'] for r in insertions],['initial','ref','quiz','exam'])
  self.assertIn('Related concept: Accuracy',insertions[1]['text']);self.assertIn(PROJECT['quiz'],insertions[2]['text'])
 def test_empty_sections_export_without_template_placeholders(self):
  d=Document(io.BytesIO(e.word({'title':'Empty'})));self.assertIn('No class materials added yet.','\n'.join(p.text for p in d.paragraphs))
  self.assertEqual(e.roman(49),'XLIX')
if __name__=='__main__':unittest.main()
