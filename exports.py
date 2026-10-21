"""Transes styling based on the user's TEMPLATE.docx, not its example content."""
import io,re,zipfile
from pathlib import Path
from lxml import etree as ET
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm,Pt,RGBColor

BLUE='1C4587';LIGHT='C9DAF8';PALE='F2F6FC'
META=[('courseCode','Course'),('professor','Professor'),('module','Module'),('semester','Semester')]

def roman(n):
    out=''
    for value,symbol in [(1000,'M'),(900,'CM'),(500,'D'),(400,'CD'),(100,'C'),(90,'XC'),(50,'L'),(40,'XL'),(10,'X'),(9,'IX'),(5,'V'),(4,'IV'),(1,'I')]:
        while n>=value:out+=symbol;n-=value
    return out

def blocks(project,kind):
    title=project.get('title') or 'Class notebook';meta=project.get('metadata') or {}
    result=[('title',title),('subtitle',{'notes':'Class notes','references':'Book references','quiz':'Quiz review','exam':'Exam review'}[kind])]
    for key,label in META:
        if meta.get(key):result.append(('meta',label+': '+str(meta[key])))
    if meta.get('student'):result.append(('meta',str(meta['student'])))
    # A section boundary keeps the heading full-width and body in two columns.
    result.append(('boundary',''))
    if kind=='notes':
        concepts=project.get('concepts',[]);result.append(('banner','TOPIC OVERVIEW'))
        for i,c in enumerate(concepts,1):result.append(('overview',roman(i)+'. '+c.get('title','Untitled concept')))
        if not concepts:result.append(('body','No class materials added yet.'))
        for i,c in enumerate(concepts,1):
            result.append(('banner',roman(i)+'. '+c.get('title','Untitled concept')))
            result.extend(('body',line) for line in c.get('text','').split('\n'))
            if c.get('source'):result.append(('source','Source: '+c['source']))
            result.append(('sideheading','SIDE NOTE/S'))
            result.extend(('side',line) for line in (c.get('side') or 'No lecture side notes added yet.').split('\n'))
    elif kind=='references':
        refs=project.get('references',[])
        for i,r in enumerate(refs,1):
            result.append(('banner',roman(i)+'. '+r.get('name','Book reference')))
            for key,label in [('location','Source location'),('related','Related concept')]:
                if r.get(key):result.append(('source',label+': '+r[key]))
            result.extend(('body',line) for line in r.get('text','').split('\n'))
        if not refs:result.append(('body','No book references added yet.'))
    else:
        result.append(('banner','IDENTIFICATION & ENUMERATION' if kind=='quiz' else 'MULTIPLE CHOICE & ELIMINATION'))
        for line in (project.get(kind) or 'No review generated yet.').split('\n'):
            style='question' if re.match(r'^\s*\d+[.)]\s',line) else 'answer' if re.match(r'^\s*(?:answer|correct answer|explanation|elimination)\b',line,re.I) else 'body'
            result.append((style,line))
    return result

def xml(parent,tag,**attrs):
    el=OxmlElement('w:'+tag)
    for key,value in attrs.items():el.set(qn('w:'+key),str(value))
    parent.append(el);return el

def page_setup(section,columns=1):
    section.page_width=Mm(210);section.page_height=Mm(297)
    section.top_margin=section.bottom_margin=section.left_margin=section.right_margin=Mm(12.7)
    cols=section._sectPr.find(qn('w:cols'))
    if cols is None:cols=xml(section._sectPr,'cols')
    cols.set(qn('w:num'),str(columns));cols.set(qn('w:space'),'720')

def shade(p,fill):xml(p._p.get_or_add_pPr(),'shd',fill=fill,val='clear')
def borders(p,color=BLUE):
    parent=xml(p._p.get_or_add_pPr(),'pBdr')
    for edge in ['top','left','bottom','right']:xml(parent,edge,val='single',sz='4',space='4',color=color)

def word(project):
    d=Document();normal=d.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(8)
    normal.paragraph_format.line_spacing=1;normal.paragraph_format.space_after=Pt(3)
    for index,kind in enumerate(['notes','references','quiz','exam']):
        section=d.sections[0] if index==0 else d.add_section(WD_SECTION_START.NEW_PAGE);page_setup(section)
        if index==0:
            footer=section.footer
            table=footer.add_table(rows=1,cols=2,width=Mm(184.6));table.columns[0].width=Mm(157);table.columns[1].width=Mm(27.6)
            for i,cell in enumerate(table.rows[0].cells):
                xml(cell._tc.get_or_add_tcPr(),'shd',fill='63AFE0' if i==0 else 'FFFFFF')
                p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                r=p.add_run(str((project.get('metadata') or {}).get('student','')) if i==0 else 'PAGE ');r.bold=True;r.font.name='Arial';r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else BLUE)
                if i==1:
                    field=p.add_run();xml(field._r,'fldChar',fldCharType='begin');xml(field._r,'instrText').text=' PAGE ';xml(field._r,'fldChar',fldCharType='end')
        for style,text in blocks(project,kind):
            if style=='meta':continue
            if style=='boundary':
                meta=project.get('metadata') or {}
                rows=[[(key,label) for key,label in META[:2]],[(key,label) for key,label in META[2:]]]
                rows=[row for row in rows if any(meta.get(key) for key,label in row)]
                if rows:
                    table=d.add_table(rows=len(rows),cols=2);table.autofit=False
                    for column in table.columns:column.width=Mm(92.3)
                    for row_index,row in enumerate(rows):
                        for cell,(key,label) in zip(table.rows[row_index].cells,row):
                            xml(cell._tc.get_or_add_tcPr(),'shd',fill=BLUE if row_index==0 else LIGHT)
                            p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                            r=p.add_run(label+': '+str(meta[key]) if meta.get(key) else '');r.bold=True;r.font.name='Poppins';r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('FFFFFF' if row_index==0 else '202020')
                page_setup(d.add_section(WD_SECTION_START.CONTINUOUS),2);continue
            p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
            r=p.add_run(text);r.font.name='Arial';r.font.size=Pt(8)
            if style in ('title','subtitle','banner','sideheading','overview'):
                r.font.name='Poppins';r.bold=True;r.font.color.rgb=RGBColor.from_string(BLUE);p.paragraph_format.keep_with_next=True
            if style=='title':r.font.name='Poppins Black';r.font.size=Pt(34);p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(3);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            elif style=='subtitle':r.font.size=Pt(11);p.paragraph_format.space_after=Pt(8)
            elif style in ('banner','sideheading'):
                shade(p,BLUE);r.font.color.rgb=RGBColor(255,255,255);r.font.size=Pt(8);p.paragraph_format.space_before=Pt(8);p.paragraph_format.space_after=Pt(4);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
            elif style=='overview':shade(p,PALE);borders(p);p.paragraph_format.keep_with_next=False
            elif style=='side':shade(p,PALE);borders(p);p.paragraph_format.space_after=Pt(3)
            elif style=='source':r.italic=True;r.font.size=Pt(7);r.font.color.rgb=RGBColor.from_string(BLUE)
            elif style in ('question','answer'):r.bold=True;r.font.color.rgb=RGBColor.from_string(BLUE)
            elif style=='meta':r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string(BLUE)
    out=io.BytesIO();d.save(out);return embed_fonts(out.getvalue())

def units(text):return len(text.encode('utf-16-le'))//2

def color(hex):return {'color':{'rgbColor':{k:int(hex[i:i+2],16)/255 for k,i in [('red',0),('green',2),('blue',4)]}}}
def dim(n):return {'magnitude':n,'unit':'PT'}

def google_plan(project,kind,tab_id):
    """Requests use Docs UTF-16 indices and explicit tab IDs throughout."""
    content=blocks(project,kind);text='';spans=[];boundary=None
    for style,line in content:
        if style=='boundary':boundary=1+units(text);continue
        start=1+units(text);text+=line+'\n';spans.append((style,start,1+units(text)))
    end=1+units(text)
    def rng(a,b):return {'tabId':tab_id,'startIndex':a,'endIndex':b}
    # Google inserts a newline plus a section break (two UTF-16 units).
    requests=[{'insertText':{'endOfSegmentLocation':{'tabId':tab_id},'text':text}}]
    if boundary is not None:
        requests.append({'insertSectionBreak':{'location':{'tabId':tab_id,'index':boundary},'sectionType':'CONTINUOUS'}})
        spans=[(style,a+2*(a>=boundary),b+2*(a>=boundary)) for style,a,b in spans];end+=2
    full=rng(1,end)
    requests += [
      {'updateDocumentStyle':{'tabId':tab_id,'documentStyle':{'pageSize':{'width':dim(595.276),'height':dim(841.89)},'marginTop':dim(36),'marginBottom':dim(36),'marginLeft':dim(36),'marginRight':dim(36)},'fields':'pageSize,marginTop,marginBottom,marginLeft,marginRight'}},
      {'updateTextStyle':{'range':full,'textStyle':{'weightedFontFamily':{'fontFamily':'Arial'},'fontSize':dim(8),'foregroundColor':color('202020')},'fields':'weightedFontFamily,fontSize,foregroundColor'}},
      {'updateParagraphStyle':{'range':full,'paragraphStyle':{'lineSpacing':100,'spaceAbove':dim(0),'spaceBelow':dim(3),'alignment':'JUSTIFIED'},'fields':'lineSpacing,spaceAbove,spaceBelow,alignment'}},
    ]
    if boundary:
        requests += [
          {'updateSectionStyle':{'range':rng(1,boundary),'sectionStyle':{'columnProperties':[{'paddingEnd':dim(0)}]},'fields':'columnProperties'}},
          {'updateSectionStyle':{'range':rng(boundary+2,end),'sectionStyle':{'columnProperties':[{'paddingEnd':dim(36)},{'paddingEnd':dim(0)}]},'fields':'columnProperties'}},
        ]
    for style,a,b in spans:
        ts={};ps={};range_=rng(a,b)
        if style in ('title','subtitle','banner','sideheading','overview'):
            ts={'weightedFontFamily':{'fontFamily':'Poppins'},'bold':True,'foregroundColor':color(BLUE)};ps={'keepWithNext':True}
        if style=='title':ts['weightedFontFamily']={'fontFamily':'Poppins','weight':900};ts['fontSize']=dim(34);ps.update(alignment='START',spaceAbove=dim(4))
        elif style=='subtitle':ts['fontSize']=dim(11);ps['spaceBelow']=dim(8)
        elif style in ('banner','sideheading'):
            ts.update(foregroundColor=color('FFFFFF'));ps.update(shading={'backgroundColor':color(BLUE)},spaceAbove=dim(8),spaceBelow=dim(4),alignment='START')
        elif style=='side':
            border={'color':color(BLUE),'width':dim(.5),'padding':dim(4),'dashStyle':'SOLID'}
            ps={'shading':{'backgroundColor':color(PALE)},**{k:dict(border) for k in ['borderTop','borderBottom','borderLeft','borderRight']}}
        elif style=='overview':ps={'shading':{'backgroundColor':color(PALE)},'keepWithNext':False}
        elif style=='source':ts={'italic':True,'fontSize':dim(7),'foregroundColor':color(BLUE)}
        elif style in ('question','answer'):ts={'bold':True,'foregroundColor':color(BLUE)}
        elif style=='meta':
            ts={'foregroundColor':color(BLUE)}

        if ts:requests.append({'updateTextStyle':{'range':range_,'textStyle':ts,'fields':','.join(ts)}})
        if ps:requests.append({'updateParagraphStyle':{'range':range_,'paragraphStyle':ps,'fields':','.join(ps)}})
    return requests

def google_export(project,token,call,read):
    doc=call('https://docs.googleapis.com/v1/documents',{'title':project.get('title') or 'Class notes'},token);did=doc['documentId'];url=f'https://docs.googleapis.com/v1/documents/{did}:batchUpdate'
    try:
        tabs=doc.get('tabs')
        if not tabs:
            import json
            tabs=json.loads(read(f'https://docs.googleapis.com/v1/documents/{did}?includeTabsContent=true',token)).get('tabs',[])
        first=tabs[0]['tabProperties']['tabId']
        titles=['Class notes','Book references','Quiz review','Exam review']
        response=call(url,{'requests':[{'updateDocumentTabProperties':{'tabProperties':{'tabId':first,'title':titles[0]},'fields':'title'}}]+[{'addDocumentTab':{'tabProperties':{'title':title}}} for title in titles[1:]]},token)
        ids=[first]+[reply['addDocumentTab']['tabProperties']['tabId'] for reply in response['replies'][1:]]
        for kind,tab in zip(['notes','references','quiz','exam'],ids):
            requests=google_plan(project,kind,tab)
            for offset in range(0,len(requests),100):call(url,{'requests':requests[offset:offset+100]},token)
    except Exception as exc:
        raise ValueError('Google document was created but formatting or content failed. Check your Drive for '+did+' before retrying. '+str(exc)) from exc
    return {'url':'https://docs.google.com/document/d/'+did+'/edit'}


def embed_fonts(data):
    """Preserve only the template's OFL Poppins families, never its personal content."""
    assets=Path(__file__).parent/'assets'/'fonts';out=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source,zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            if entry.filename in ('word/fontTable.xml','word/_rels/fontTable.xml.rels'):continue
            content=source.read(entry.filename)
            if entry.filename=='[Content_Types].xml':
                root=ET.fromstring(content);ns='{http://schemas.openxmlformats.org/package/2006/content-types}'
                ET.SubElement(root,ns+'Default',{'Extension':'ttf','ContentType':'application/x-font-ttf'})
                content=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            target.writestr(entry,content)
        for file in assets.glob('*.ttf'):target.writestr('word/fonts/'+file.name,file.read_bytes())
        target.writestr('word/fontTable.xml',(assets/'fontTable.xml').read_bytes())
        target.writestr('word/_rels/fontTable.xml.rels',(assets/'fontTable.xml.rels').read_bytes())
    return out.getvalue()
